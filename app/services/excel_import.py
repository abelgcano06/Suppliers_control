"""Carga y validacion del reporte de ordenes abiertas (Excel del sistema de compras).

Lo que hace una carga, en orden:

1. Encuentra la fila de encabezados y mapea las columnas a nombres canonicos.
2. Valida fila por fila. Una fila invalida se rechaza y se reporta; no tumba la carga.
3. Da de alta los proveedores que aparezcan por primera vez.
4. Inserta o actualiza cada linea de orden, sobreescribiendo SOLO los campos que
   controla el sistema. El status, la fecha promesa y el comentario del proveedor
   nunca se tocan desde aqui.
5. Cierra las ordenes abiertas que ya no vienen en el reporte.

Cada diferencia se escribe en `order_changes`, asi que el historial queda completo.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    ChangeSource,
    OrderChange,
    OrderStatus,
    PurchaseOrder,
    Supplier,
    UploadBatch,
    utcnow,
)

# Encabezados aceptados para cada campo. Se comparan normalizados (sin acentos,
# minusculas, sin espacios ni signos), asi que "Fecha Requerida" == "fecha_requerida".
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "po_number": ("orden", "ordendecompra", "oc", "po", "ponumber", "numeroorden", "purchaseorder", "documento"),
    "line_number": ("linea", "numerolinea", "line", "linenumber", "partida", "posicion", "item"),
    "supplier_code": ("codigoproveedor", "proveedorcodigo", "suppliercode", "vendor", "vendorcode", "numeroproveedor", "idproveedor"),
    "supplier_name": ("proveedor", "nombreproveedor", "suppliername", "vendorname", "razonsocial"),
    "part_number": ("parte", "numeroparte", "partnumber", "numparte", "material", "sku", "codigoparte"),
    "description": ("descripcion", "description", "detalle", "concepto", "texto"),
    "quantity": ("cantidad", "quantity", "qty", "cant", "cantidadpendiente"),
    "unit": ("unidad", "um", "unit", "uom", "unidadmedida"),
    "unit_price": ("precio", "preciounitario", "unitprice", "price", "costo", "costounitario"),
    "currency": ("moneda", "currency", "divisa"),
    "required_date": ("fecharequerida", "fecharequerimiento", "requireddate", "fechaentrega", "fechaentregarequerida", "duedate", "needdate"),
}

REQUIRED_COLUMNS = ("po_number", "line_number")
# Campos que controla el sistema: la carga los sobreescribe siempre.
SYSTEM_FIELDS = (
    "part_number",
    "description",
    "quantity",
    "unit",
    "unit_price",
    "currency",
    "required_date",
)

MAX_HEADER_SCAN_ROWS = 15


class ExcelImportError(Exception):
    """El archivo no se puede procesar (formato, hoja vacia, faltan columnas)."""


@dataclass
class ImportResult:
    batch_id: int
    rows_total: int = 0
    rows_created: int = 0
    rows_updated: int = 0
    rows_unchanged: int = 0
    rows_rejected: int = 0
    orders_closed: int = 0
    suppliers_created: int = 0
    errors: list[str] = field(default_factory=list)
    # Columnas del archivo que no se reconocieron (informativo).
    ignored_columns: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.rows_rejected == 0


# --------------------------------------------------------------------------- #
# Normalizacion y parseo de valores
# --------------------------------------------------------------------------- #

_ACCENTS = str.maketrans("áàäâéèëêíìïîóòöôúùüûñ", "aaaaeeeeiiiioooouuuun")


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower().translate(_ACCENTS)
    return "".join(ch for ch in text if ch.isalnum())


def parse_text(value: object, max_length: int | None = None) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if max_length and len(text) > max_length:
        text = text[:max_length]
    return text


def parse_number(value: object) -> Decimal | None:
    """Acepta 1234.56, '1,234.56', '$1,234.56', '(100)' como negativo."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "").replace("$", "").replace(" ", "")
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        raise ValueError(f"'{value}' no es un numero valido")
    return -number if negative else number


def parse_int(value: object) -> int | None:
    number = parse_number(value)
    if number is None:
        return None
    if number != number.to_integral_value():
        raise ValueError(f"'{value}' debe ser un numero entero")
    return int(number)


_DATE_FORMATS = (
    "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d",
    "%d.%m.%Y", "%Y%m%d", "%d/%m/%y", "%m/%d/%y",
)


def parse_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Numero de serie de Excel (base 1899-12-30).
        from datetime import timedelta

        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{value}' no es una fecha valida (usa AAAA-MM-DD o DD/MM/AAAA)")


def _as_display(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, OrderStatus):
        return value.value
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value.normalize(), "f")
    return str(value)


# --------------------------------------------------------------------------- #
# Lectura del archivo
# --------------------------------------------------------------------------- #


@dataclass
class ParsedSheet:
    column_map: dict[str, int]
    ignored_columns: list[str]
    rows: list[tuple[int, list]]  # (numero de fila en Excel, celdas)


def read_sheet(content: bytes) -> ParsedSheet:
    """Abre el Excel, localiza el encabezado y devuelve las filas con datos."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza varios tipos segun el daño
        raise ExcelImportError(
            "No se pudo leer el archivo. Debe ser un Excel .xlsx valido "
            f"(detalle: {exc})."
        ) from exc

    sheet = workbook.active
    if sheet is None:
        raise ExcelImportError("El archivo no tiene hojas.")

    all_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()

    header_index = _find_header_row(all_rows)
    if header_index is None:
        faltantes = ", ".join(REQUIRED_COLUMNS)
        raise ExcelImportError(
            "No se encontro la fila de encabezados. El reporte debe traer al menos "
            f"las columnas: {faltantes} (por ejemplo 'Orden' y 'Linea')."
        )

    header = all_rows[header_index]
    column_map, ignored = _map_columns(header)

    missing = [c for c in REQUIRED_COLUMNS if c not in column_map]
    if missing:
        raise ExcelImportError(
            "Al reporte le faltan columnas obligatorias: " + ", ".join(missing)
        )

    rows: list[tuple[int, list]] = []
    for offset, cells in enumerate(all_rows[header_index + 1 :], start=header_index + 2):
        if all(cell is None or str(cell).strip() == "" for cell in cells):
            continue  # fila vacia
        rows.append((offset, cells))

    return ParsedSheet(column_map=column_map, ignored_columns=ignored, rows=rows)


def _find_header_row(rows: list[list]) -> int | None:
    for index, cells in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
        column_map, _ = _map_columns(cells)
        if all(c in column_map for c in REQUIRED_COLUMNS):
            return index
    return None


def _map_columns(header: list) -> tuple[dict[str, int], list[str]]:
    column_map: dict[str, int] = {}
    ignored: list[str] = []
    for position, raw in enumerate(header):
        key = normalize_header(raw)
        if not key:
            continue
        for canonical, aliases in COLUMN_ALIASES.items():
            if canonical in column_map:
                continue  # la primera columna que coincide gana
            if key in aliases or key == canonical:
                column_map[canonical] = position
                break
        else:
            ignored.append(str(raw).strip())
    return column_map, ignored


def _cell(cells: list, column_map: dict[str, int], name: str):
    position = column_map.get(name)
    if position is None or position >= len(cells):
        return None
    return cells[position]


# --------------------------------------------------------------------------- #
# Importacion
# --------------------------------------------------------------------------- #


def import_orders(
    db: Session,
    content: bytes,
    filename: str,
    uploaded_by: str | None = None,
    close_scope: str = "suppliers_in_file",
    create_missing_suppliers: bool = True,
) -> ImportResult:
    """Procesa el reporte y deja la base de datos al dia.

    close_scope:
      * "all"               - el archivo trae TODAS las ordenes abiertas; cierra
                              cualquier orden abierta que no aparezca.
      * "suppliers_in_file" - cierra solo dentro de los proveedores presentes en
                              el archivo (por defecto, seguro para cargas parciales).
      * "none"              - no cierra nada.
    """
    if close_scope not in ("all", "suppliers_in_file", "none"):
        raise ValueError(f"close_scope invalido: {close_scope}")

    sheet = read_sheet(content)

    batch = UploadBatch(filename=filename, uploaded_by=uploaded_by)
    db.add(batch)
    db.flush()  # necesitamos batch.id para marcar las ordenes

    result = ImportResult(batch_id=batch.id, ignored_columns=sheet.ignored_columns)

    suppliers_by_code = {s.code: s for s in db.scalars(select(Supplier)).all()}
    seen_keys: set[tuple[str, int]] = set()
    seen_supplier_ids: set[int] = set()

    for excel_row, cells in sheet.rows:
        result.rows_total += 1
        try:
            parsed = _parse_row(cells, sheet.column_map)
        except ValueError as exc:
            result.rows_rejected += 1
            result.errors.append(f"Fila {excel_row}: {exc}")
            continue

        key = (parsed["po_number"], parsed["line_number"])
        if key in seen_keys:
            result.rows_rejected += 1
            result.errors.append(
                f"Fila {excel_row}: la linea {key[0]}-{key[1]} viene repetida en el archivo."
            )
            continue

        supplier = _resolve_supplier(
            db, suppliers_by_code, parsed, create_missing_suppliers
        )
        if supplier is None:
            result.rows_rejected += 1
            code = parsed["supplier_code"] or parsed["supplier_name"] or "(sin dato)"
            result.errors.append(
                f"Fila {excel_row}: el proveedor '{code}' no esta dado de alta en el portal."
            )
            continue
        if supplier.id is None:
            db.flush()
            result.suppliers_created += 1

        seen_keys.add(key)
        seen_supplier_ids.add(supplier.id)

        order = db.scalar(
            select(PurchaseOrder).where(
                PurchaseOrder.po_number == key[0],
                PurchaseOrder.line_number == key[1],
            )
        )
        if order is None:
            _create_order(db, parsed, supplier, batch)
            result.rows_created += 1
        else:
            changed = _update_order(db, order, parsed, supplier, batch)
            if changed:
                result.rows_updated += 1
            else:
                result.rows_unchanged += 1

    if close_scope != "none":
        result.orders_closed = _close_missing_orders(
            db, seen_keys, seen_supplier_ids, close_scope, batch
        )

    batch.rows_total = result.rows_total
    batch.rows_created = result.rows_created
    batch.rows_updated = result.rows_updated
    batch.rows_unchanged = result.rows_unchanged
    batch.rows_rejected = result.rows_rejected
    batch.orders_closed = result.orders_closed
    batch.suppliers_created = result.suppliers_created
    batch.error_log = "\n".join(result.errors) if result.errors else None

    db.commit()
    return result


def _parse_row(cells: list, column_map: dict[str, int]) -> dict:
    po_number = parse_text(_cell(cells, column_map, "po_number"), 50)
    if not po_number:
        raise ValueError("falta el numero de orden")

    raw_line = _cell(cells, column_map, "line_number")
    line_number = parse_int(raw_line)
    if line_number is None:
        raise ValueError("falta el numero de linea")
    if line_number < 0:
        raise ValueError("el numero de linea no puede ser negativo")

    quantity = parse_number(_cell(cells, column_map, "quantity"))
    if quantity is not None and quantity < 0:
        raise ValueError("la cantidad no puede ser negativa")

    unit_price = parse_number(_cell(cells, column_map, "unit_price"))
    if unit_price is not None and unit_price < 0:
        raise ValueError("el precio no puede ser negativo")

    return {
        "po_number": po_number,
        "line_number": line_number,
        "supplier_code": parse_text(_cell(cells, column_map, "supplier_code"), 50),
        "supplier_name": parse_text(_cell(cells, column_map, "supplier_name"), 200),
        "part_number": parse_text(_cell(cells, column_map, "part_number"), 100),
        "description": parse_text(_cell(cells, column_map, "description"), 500),
        "quantity": quantity,
        "unit": parse_text(_cell(cells, column_map, "unit"), 20),
        "unit_price": unit_price,
        "currency": parse_text(_cell(cells, column_map, "currency"), 10),
        "required_date": parse_date(_cell(cells, column_map, "required_date")),
    }


def _resolve_supplier(
    db: Session,
    cache: dict[str, Supplier],
    parsed: dict,
    create_missing: bool,
) -> Supplier | None:
    # El codigo manda; si el reporte no lo trae, se usa el nombre como codigo.
    code = parsed["supplier_code"] or parsed["supplier_name"]
    if not code:
        return None
    supplier = cache.get(code)
    if supplier is not None:
        # Refrescar el nombre si el reporte trae uno mejor.
        if parsed["supplier_name"] and supplier.name != parsed["supplier_name"]:
            supplier.name = parsed["supplier_name"]
        return supplier
    if not create_missing:
        return None
    supplier = Supplier(code=code, name=parsed["supplier_name"] or code)
    db.add(supplier)
    cache[code] = supplier
    return supplier


def _create_order(db: Session, parsed: dict, supplier: Supplier, batch: UploadBatch) -> PurchaseOrder:
    order = PurchaseOrder(
        po_number=parsed["po_number"],
        line_number=parsed["line_number"],
        supplier_id=supplier.id,
        status=OrderStatus.PENDIENTE,
        first_seen_batch_id=batch.id,
        last_seen_batch_id=batch.id,
    )
    for name in SYSTEM_FIELDS:
        setattr(order, name, parsed[name])
    db.add(order)
    db.flush()
    db.add(
        OrderChange(
            order_id=order.id,
            field="orden",
            old_value=None,
            new_value="alta",
            source=ChangeSource.EXCEL,
            actor=batch.uploaded_by,
        )
    )
    return order


def _update_order(
    db: Session, order: PurchaseOrder, parsed: dict, supplier: Supplier, batch: UploadBatch
) -> bool:
    """Sobreescribe los campos del sistema. Devuelve True si algo cambio."""
    changed = False

    if order.supplier_id != supplier.id:
        db.add(
            OrderChange(
                order_id=order.id,
                field="supplier_id",
                old_value=str(order.supplier_id),
                new_value=str(supplier.id),
                source=ChangeSource.EXCEL,
                actor=batch.uploaded_by,
            )
        )
        order.supplier_id = supplier.id
        changed = True

    for name in SYSTEM_FIELDS:
        nuevo = parsed[name]
        actual = getattr(order, name)
        if _same_value(actual, nuevo):
            continue
        db.add(
            OrderChange(
                order_id=order.id,
                field=name,
                old_value=_as_display(actual),
                new_value=_as_display(nuevo),
                source=ChangeSource.EXCEL,
                actor=batch.uploaded_by,
            )
        )
        setattr(order, name, nuevo)
        changed = True

    if order.is_closed:
        # La orden volvio a aparecer en el reporte: se reabre.
        db.add(
            OrderChange(
                order_id=order.id,
                field="is_closed",
                old_value="Cerrada",
                new_value="Abierta",
                source=ChangeSource.EXCEL,
                actor=batch.uploaded_by,
            )
        )
        order.is_closed = False
        order.closed_at = None
        changed = True

    order.last_seen_batch_id = batch.id
    if changed:
        order.updated_at = utcnow()
    return changed


def _same_value(actual, nuevo) -> bool:
    if actual is None and nuevo is None:
        return True
    if actual is None or nuevo is None:
        return False
    if isinstance(actual, Decimal) or isinstance(nuevo, Decimal):
        try:
            return Decimal(str(actual)) == Decimal(str(nuevo))
        except (InvalidOperation, ValueError):
            return str(actual) == str(nuevo)
    return actual == nuevo


def _close_missing_orders(
    db: Session,
    seen_keys: set[tuple[str, int]],
    seen_supplier_ids: set[int],
    close_scope: str,
    batch: UploadBatch,
) -> int:
    query = select(PurchaseOrder).where(PurchaseOrder.is_closed == False)  # noqa: E712
    if close_scope == "suppliers_in_file":
        if not seen_supplier_ids:
            return 0
        query = query.where(PurchaseOrder.supplier_id.in_(seen_supplier_ids))

    closed = 0
    for order in db.scalars(query).all():
        if (order.po_number, order.line_number) in seen_keys:
            continue
        order.is_closed = True
        order.closed_at = utcnow()
        db.add(
            OrderChange(
                order_id=order.id,
                field="is_closed",
                old_value="Abierta",
                new_value="Cerrada",
                source=ChangeSource.EXCEL,
                actor=batch.uploaded_by,
            )
        )
        closed += 1
    return closed
