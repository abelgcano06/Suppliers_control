"""Portal interno: las pantallas que usa el equipo de mantenimiento."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import desc, select
from sqlalchemy.orm import Session, selectinload

from app.config import get_settings
from app.db import get_db
from app.models import (
    ChangeSource,
    OrderChange,
    OrderStatus,
    PurchaseOrder,
    Supplier,
    UploadBatch,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services import dashboard, seguimiento, sharepoint
from app.services.excel_import import ExcelImportError, import_orders

router = APIRouter(tags=["portal"])
templates = Jinja2Templates(directory="app/templates")

MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB


def _context(**extra) -> dict:
    settings = get_settings()
    base = {
        "settings": settings,
        "hoy": date.today(),
        "OrderStatus": OrderStatus,
    }
    base.update(extra)
    return base


# --------------------------------------------------------------------------- #
# Tablero
# --------------------------------------------------------------------------- #


@router.get("/", response_class=HTMLResponse)
def tablero(
    request: Request,
    proveedor: int | None = None,
    status_filtro: str | None = None,
    alerta: str | None = None,
    q: str | None = None,
    cerradas: bool = False,
    db: Session = Depends(get_db),
):
    filtros = dashboard.OrderFilters(
        supplier_id=proveedor,
        status=_status_o_none(status_filtro),
        alerta=alerta or None,
        search=q,
        include_closed=cerradas,
    )
    ordenes = dashboard.query_orders(db, filtros)
    return templates.TemplateResponse(
        request,
        "tablero.html",
        _context(
            ordenes=ordenes,
            kpis=dashboard.compute_kpis(db),
            horas_respuesta=dashboard.tiempo_promedio_respuesta_horas(db),
            proveedores=db.scalars(select(Supplier).order_by(Supplier.name)).all(),
            filtros=filtros,
            filtro_status=status_filtro or "",
            filtro_alerta=alerta or "",
            busqueda=q or "",
            ultima_carga=db.scalar(select(UploadBatch).order_by(desc(UploadBatch.uploaded_at))),
        ),
    )


def _status_o_none(valor: str | None) -> OrderStatus | None:
    if not valor:
        return None
    for s in OrderStatus:
        if valor == s.name or valor == s.value:
            return s
    return None


@router.get("/ordenes/{order_id}", response_class=HTMLResponse)
def detalle_orden(order_id: int, request: Request, db: Session = Depends(get_db)):
    orden = db.scalar(
        select(PurchaseOrder)
        .options(selectinload(PurchaseOrder.supplier), selectinload(PurchaseOrder.changes))
        .where(PurchaseOrder.id == order_id)
    )
    if orden is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Orden no encontrada.")
    return templates.TemplateResponse(request, "orden.html", _context(orden=orden))


@router.post("/ordenes/{order_id}/seguimiento")
def guardar_seguimiento(
    order_id: int,
    fecha_promesa: str = Form(""),
    nota: str = Form(""),
    capturado_por: str = Form(""),
    db: Session = Depends(get_db),
):
    """Fecha promesa y nota interna: las captura mantenimiento, no el proveedor."""
    orden = db.get(PurchaseOrder, order_id)
    if orden is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Orden no encontrada.")

    texto = fecha_promesa.strip()
    promesa = None
    if texto:
        try:
            promesa = date.fromisoformat(texto)
        except ValueError:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                f"'{texto}' no es una fecha valida. Usa el selector de fecha.",
            )

    seguimiento.actualizar_seguimiento(
        db,
        orden,
        promised_date=promesa,
        internal_note=nota,
        actor=capturado_por.strip() or None,
        limpiar_fecha=not texto,
    )
    db.commit()
    return RedirectResponse(f"/ordenes/{order_id}", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/export.csv")
def exportar_csv(
    proveedor: int | None = None,
    status_filtro: str | None = None,
    alerta: str | None = None,
    q: str | None = None,
    cerradas: bool = False,
    db: Session = Depends(get_db),
):
    """El tablero tal como se ve, en CSV, para pegarlo en un reporte."""
    filtros = dashboard.OrderFilters(
        supplier_id=proveedor,
        status=_status_o_none(status_filtro),
        alerta=alerta or None,
        search=q,
        include_closed=cerradas,
    )
    ordenes = dashboard.query_orders(db, filtros)

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "Orden", "Linea", "Proveedor", "Parte", "Descripcion", "Cantidad", "Unidad",
            "Fecha requerida", "Status", "Fecha promesa", "Dias de retraso",
            "Nota interna", "Ultima respuesta", "Cerrada",
        ]
    )
    hoy = date.today()
    for o in ordenes:
        writer.writerow(
            [
                o.po_number, o.line_number, o.supplier.name if o.supplier else "",
                o.part_number or "", o.description or "",
                float(o.quantity) if o.quantity is not None else "", o.unit or "",
                o.required_date.isoformat() if o.required_date else "",
                o.status.value,
                o.promised_date.isoformat() if o.promised_date else "",
                o.dias_de_retraso(hoy) if o.esta_retrasada(hoy) else 0,
                o.internal_note or "",
                o.last_supplier_update_at.isoformat(sep=" ", timespec="minutes")
                if o.last_supplier_update_at else "",
                "Si" if o.is_closed else "No",
            ]
        )

    buffer.seek(0)
    nombre = f"ordenes-{hoy.isoformat()}.csv"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


# --------------------------------------------------------------------------- #
# Carga del reporte
# --------------------------------------------------------------------------- #


@router.get("/cargar", response_class=HTMLResponse)
def pantalla_carga(request: Request, db: Session = Depends(get_db)):
    cargas = db.scalars(
        select(UploadBatch).order_by(desc(UploadBatch.uploaded_at)).limit(15)
    ).all()
    return templates.TemplateResponse(
        request, "cargar.html", _context(cargas=cargas, resultado=None, error=None)
    )


@router.post("/cargar", response_class=HTMLResponse)
async def procesar_carga(
    request: Request,
    archivo: UploadFile = File(...),
    cargado_por: str = Form(""),
    alcance_cierre: str = Form("suppliers_in_file"),
    alta_automatica: bool = Form(False),
    db: Session = Depends(get_db),
):
    contenido = await archivo.read()
    resultado = None
    error = None

    if alcance_cierre not in ("all", "suppliers_in_file", "none"):
        alcance_cierre = "suppliers_in_file"

    if not contenido:
        error = "El archivo llego vacio."
    elif len(contenido) > MAX_UPLOAD_BYTES:
        error = f"El archivo pesa mas de {MAX_UPLOAD_BYTES // (1024 * 1024)} MB."
    elif not (archivo.filename or "").lower().endswith((".xlsx", ".xlsm")):
        error = "Solo se aceptan archivos .xlsx o .xlsm. Guarda el reporte en ese formato."
    else:
        try:
            resultado = import_orders(
                db,
                content=contenido,
                filename=archivo.filename or "reporte.xlsx",
                uploaded_by=cargado_por.strip() or None,
                close_scope=alcance_cierre,
                create_missing_suppliers=alta_automatica,
            )
        except ExcelImportError as exc:
            db.rollback()
            error = str(exc)

    cargas = db.scalars(
        select(UploadBatch).order_by(desc(UploadBatch.uploaded_at)).limit(15)
    ).all()
    return templates.TemplateResponse(
        request,
        "cargar.html",
        _context(cargas=cargas, resultado=resultado, error=error),
    )


@router.get("/plantilla.csv")
def plantilla():
    """Plantilla con los encabezados que el portal espera del reporte."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        ["Orden", "Linea", "Codigo Proveedor", "Proveedor", "Parte", "Descripcion",
         "Cantidad", "Unidad", "Precio Unitario", "Moneda", "Fecha Requerida"]
    )
    writer.writerow(
        ["OC-1001", "1", "PROV-A", "Pinturas del Norte", "PN-4471",
         "Filtro de cabina de pintura", "12", "PZA", "845.50", "MXN", "2026-10-15"]
    )
    buffer.seek(0)
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="plantilla-ordenes.csv"'},
    )


# --------------------------------------------------------------------------- #
# Proveedores
# --------------------------------------------------------------------------- #


@router.get("/proveedores", response_class=HTMLResponse)
def pantalla_proveedores(request: Request, db: Session = Depends(get_db)):
    return templates.TemplateResponse(
        request,
        "proveedores.html",
        _context(resumen=dashboard.supplier_summaries(db), error=None),
    )


@router.post("/proveedores")
def alta_proveedor(
    request: Request,
    codigo: str = Form(...),
    nombre: str = Form(...),
    correo: str = Form(""),
    lista_sharepoint: str = Form(""),
    compartir_precio: bool = Form(False),
    db: Session = Depends(get_db),
):
    codigo = codigo.strip()
    error = None
    if not codigo or not nombre.strip():
        error = "El codigo y el nombre son obligatorios."
    elif db.scalar(select(Supplier).where(Supplier.code == codigo)):
        error = f"Ya existe un proveedor con el codigo '{codigo}'. Usa otro o edita el existente."

    if error:
        # Se regresa la pantalla con el mensaje, no un error crudo.
        return templates.TemplateResponse(
            request,
            "proveedores.html",
            _context(resumen=dashboard.supplier_summaries(db), error=error),
            status_code=status.HTTP_409_CONFLICT,
        )

    db.add(
        Supplier(
            code=codigo,
            name=nombre.strip(),
            contact_email=correo.strip() or None,
            sharepoint_list=lista_sharepoint.strip() or None,
            share_price=compartir_precio,
        )
    )
    db.commit()
    return RedirectResponse("/proveedores", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/proveedores/{supplier_id}")
def editar_proveedor(
    supplier_id: int,
    nombre: str = Form(...),
    correo: str = Form(""),
    lista_sharepoint: str = Form(""),
    compartir_precio: bool = Form(False),
    activo: bool = Form(False),
    db: Session = Depends(get_db),
):
    proveedor = db.get(Supplier, supplier_id)
    if proveedor is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Proveedor no encontrado.")
    proveedor.name = nombre.strip() or proveedor.name
    proveedor.contact_email = correo.strip() or None
    proveedor.sharepoint_list = lista_sharepoint.strip() or None
    proveedor.share_price = compartir_precio
    proveedor.active = activo
    db.commit()
    return RedirectResponse("/proveedores", status_code=status.HTTP_303_SEE_OTHER)


# --------------------------------------------------------------------------- #
# Exportar un archivo por proveedor
# --------------------------------------------------------------------------- #

# Caracteres que Windows no acepta en un nombre de archivo.
_PROHIBIDOS = re.compile(r'[<>:"/\\|?*]')


def _nombre_de_archivo(supplier: Supplier) -> str:
    limpio = _PROHIBIDOS.sub("-", f"{supplier.code} - {supplier.name}").strip()
    return f"{limpio[:80]}.xlsx"


def _excel_del_proveedor(supplier: Supplier, ordenes: list[PurchaseOrder]) -> bytes:
    """Un Excel listo para subir a la lista del proveedor, o para mandarselo por correo."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ordenes abiertas"

    encabezados = ["Clave", "Orden", "Linea", "Parte", "Descripcion", "Cantidad", "Unidad",
                   "FechaRequerida"]
    if supplier.share_price:
        encabezados += ["Precio", "Moneda"]
    # Status es lo unico que el proveedor captura.
    encabezados += ["Status"]

    hoja.append(encabezados)
    relleno = PatternFill("solid", fgColor="1F4E79")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")

    for o in ordenes:
        fila = [
            f"{o.po_number}-{o.line_number}", o.po_number, o.line_number,
            o.part_number or "", o.description or "",
            float(o.quantity) if o.quantity is not None else None, o.unit or "",
            o.required_date,
        ]
        if supplier.share_price:
            fila += [float(o.unit_price) if o.unit_price is not None else None, o.currency or ""]
        fila += [o.status.value]
        hoja.append(fila)

    for columna, ancho in zip("ABCDEFGHIJK", (16, 12, 6, 14, 38, 10, 8, 15, 12, 8, 14)):
        hoja.column_dimensions[columna].width = ancho
    hoja.freeze_panes = "A2"

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


@router.get("/proveedores/exportar.zip")
def exportar_por_proveedor(db: Session = Depends(get_db)):
    """Un Excel por proveedor, ya clasificado, en un solo ZIP.

    Es el camino que no depende de permisos de nada: si el tenant no deja que el
    portal escriba en SharePoint, subes estos archivos a mano.
    """
    proveedores = db.scalars(
        select(Supplier).where(Supplier.active == True).order_by(Supplier.name)  # noqa: E712
    ).all()

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zip_archivo:
        for proveedor in proveedores:
            ordenes = dashboard.query_orders(
                db, dashboard.OrderFilters(supplier_id=proveedor.id)
            )
            if not ordenes:
                continue
            zip_archivo.writestr(
                _nombre_de_archivo(proveedor), _excel_del_proveedor(proveedor, ordenes)
            )

    buffer.seek(0)
    nombre = f"ordenes-por-proveedor-{date.today().isoformat()}.zip"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


# --------------------------------------------------------------------------- #
# Historial de cambios
# --------------------------------------------------------------------------- #


@router.get("/historial", response_class=HTMLResponse)
def historial(request: Request, origen: str | None = None, db: Session = Depends(get_db)):
    query = (
        select(OrderChange)
        .options(selectinload(OrderChange.order).selectinload(PurchaseOrder.supplier))
        .order_by(desc(OrderChange.changed_at))
        .limit(300)
    )
    if origen:
        for fuente in ChangeSource:
            if origen == fuente.name:
                query = query.where(OrderChange.source == fuente)
                break
    return templates.TemplateResponse(
        request,
        "historial.html",
        _context(
            cambios=db.scalars(query).all(),
            ChangeSource=ChangeSource,
            filtro_origen=origen or "",
        ),
    )


# --------------------------------------------------------------------------- #
# SharePoint
# --------------------------------------------------------------------------- #


def _pantalla_sharepoint(request: Request, db: Session, resultado=None, error=None,
                         codigo_http: int = 200):
    settings = get_settings()
    try:
        cuenta = sharepoint.sesion_activa(settings)
    except sharepoint.SharePointError as exc:
        cuenta, error = None, error or str(exc)

    proveedores = dashboard.supplier_summaries(db)
    return templates.TemplateResponse(
        request,
        "sharepoint.html",
        _context(
            cuenta=cuenta,
            sitio=settings.sp_site_url,
            modo=settings.sp_auth_mode,
            resultado=resultado,
            error=error,
            sin_lista=[r.supplier for r in proveedores
                       if r.supplier.active and not r.supplier.sharepoint_list],
            con_lista=[r for r in proveedores
                       if r.supplier.active and r.supplier.sharepoint_list],
        ),
        status_code=codigo_http,
    )


@router.get("/sharepoint", response_class=HTMLResponse)
def pantalla_sharepoint(request: Request, db: Session = Depends(get_db)):
    return _pantalla_sharepoint(request, db)


@router.post("/sharepoint/sincronizar", response_class=HTMLResponse)
def sincronizar_sharepoint(request: Request, db: Session = Depends(get_db)):
    """Baja lo que capturaron los proveedores y publica las ordenes al dia."""
    settings = get_settings()
    if not settings.sp_site_url:
        return _pantalla_sharepoint(
            request, db,
            error="Falta capturar SP_SITE_URL en el archivo .env (la liga de tu sitio).",
            codigo_http=400,
        )
    try:
        resultado = sharepoint.sincronizar(db, settings)
    except sharepoint.SinSesion as exc:
        return _pantalla_sharepoint(request, db, error=str(exc), codigo_http=401)
    except sharepoint.SharePointError as exc:
        db.rollback()
        return _pantalla_sharepoint(request, db, error=str(exc), codigo_http=502)
    return _pantalla_sharepoint(request, db, resultado=resultado)


@router.post("/sharepoint/desconectar")
def desconectar_sharepoint():
    sharepoint.cerrar_sesion()
    return RedirectResponse("/sharepoint", status_code=status.HTTP_303_SEE_OTHER)
